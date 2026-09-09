# -*- coding: utf-8 -*-
"""第8组交付文档：按当前代码状态重写规约/矩阵/用例/说明书/实习报告。"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor, Inches
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
TODAY = "2026-09-09"
VER = "V2.0"
GROUP, LEADER, TITLE = "第8组", "杜青桐", "智能制造云平台"
MEMBERS = "杜青桐、李雨桐、马嘉祺、肖舒予"
OWNERS = ["杜青桐", "李雨桐", "马嘉祺", "肖舒予"]

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
THIN = Border(
    left=Side(style="thin", color="B0B8C4"),
    right=Side(style="thin", color="B0B8C4"),
    top=Side(style="thin", color="B0B8C4"),
    bottom=Side(style="thin", color="B0B8C4"),
)
HFILL = PatternFill("solid", fgColor="1F3A5F")
HFONT = Font(name="微软雅黑", bold=True, color="FFFFFF", size=9)
CFONT = Font(name="微软雅黑", size=9)
WRAP = Alignment(wrap_text=True, vertical="center")
OKFILL = PatternFill("solid", fgColor="C6EFCE")
YFILL = PatternFill("solid", fgColor="FFF2CC")
RFILL = PatternFill("solid", fgColor="FCE4D6")


def set_cn(run, name="宋体", size=12, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color


def add_p(doc, text, size=12, bold=False, align="left", after=6, name="宋体", first=0, color=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(after)
    pf.space_before = Pt(0)
    pf.line_spacing = 1.25
    if first:
        pf.first_line_indent = Cm(first)
    p.alignment = {
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "both": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }.get(align, WD_ALIGN_PARAGRAPH.LEFT)
    run = p.add_run(text)
    set_cn(run, name, size, bold, color)
    return p


def add_h(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.color.rgb = NAVY
        run.font.name = "黑体"
        rPr = run._element.get_or_add_rPr()
        rFonts = rPr.get_or_add_rFonts()
        rFonts.set(qn("w:eastAsia"), "黑体")
    return p


def shade_header(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    from docx.oxml import OxmlElement
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "1F3A5F")
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def add_table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        shade_header(cell)
        r = cell.paragraphs[0].add_run(h)
        set_cn(r, "黑体", 9, True, RGBColor(0xFF, 0xFF, 0xFF))
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = t.rows[ri + 1].cells[ci]
            cell.text = ""
            r = cell.paragraphs[0].add_run(str(val))
            set_cn(r, "宋体", 9)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                if i < len(row.cells):
                    row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return t


def add_flow(doc, title, lines):
    add_p(doc, title, 11, True, after=4)
    box = "\n".join(lines)
    p = add_p(doc, box, 9, False, after=10, name="等线")
    p.paragraph_format.line_spacing = 1.15


def setup_doc(footer_name):
    doc = Document()
    s = doc.sections[0]
    s.top_margin = Cm(2.2)
    s.bottom_margin = Cm(2.2)
    s.left_margin = Cm(2.5)
    s.right_margin = Cm(2.2)
    footer = s.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run(f"{TITLE}  {GROUP}  {LEADER}  {VER}  {TODAY}  密级：内部")
    set_cn(r, "宋体", 9, color=RGBColor(0x66, 0x66, 0x66))
    return doc


def cover(doc, doc_name, extra=""):
    add_p(doc, "东软集团股份有限公司  IT 人才实训中心", 12, True, "center", 4)
    add_p(doc, "项目交付文档", 14, False, "center", 18)
    add_p(doc, TITLE, 26, True, "center", 8, "黑体")
    add_p(doc, doc_name, 18, True, "center", 12, "黑体")
    add_p(doc, extra, 12, False, "center", 8)
    add_table(doc, ["项", "内容"], [
        ["项目名称", TITLE + "（制造协同与履约平台）"],
        ["文档版本", VER],
        ["编制日期", TODAY],
        ["小组", f"{GROUP}    {MEMBERS}"],
        ["编制 / 审核", f"{LEADER} 编制；小组成员交叉审核"],
        ["适用系统", "Vue 3 + Spring Boot 3 + MySQL 8  本地演示环境"],
        ["密级", "内部 · 实训提交"],
    ], [4, 13])


def write_rows(ws, headers, rows, widths, header_row=1):
    for c, h in enumerate(headers, 1):
        cell = ws.cell(header_row, c, h)
        cell.fill = HFILL
        cell.font = HFONT
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = THIN
    for i, row in enumerate(rows):
        for c, v in enumerate(row, 1):
            cell = ws.cell(header_row + 1 + i, c, v)
            cell.font = CFONT
            cell.alignment = WRAP
            cell.border = THIN
            if len(row) >= 4 and str(row[3]) in ("异常", "安全", "边界"):
                if c == 4:
                    cell.fill = RFILL if row[3] != "边界" else YFILL
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[header_row].height = 24
    ws.freeze_panes = f"A{header_row + 1}"
    if rows:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(headers))}{header_row + len(rows)}"
    for r in range(header_row + 1, header_row + 1 + len(rows)):
        ws.row_dimensions[r].height = 56


# ========================= 测试数据（演示库可对照） =========================
DEMO = """买家 13000000001 / 123456  杭州精工传动有限公司  信用代码 91330110MA2K8B1X1A  信用分 82
工厂 13000000002 宁波博锐精密机械有限公司
工厂 13000000003 台州宏达机械加工厂
工厂 13000000004 嘉兴金盾热处理有限公司
工厂 13000000005 苏州汇通智能制造有限公司
运营 admin / admin123    质检 inspect01 / 123456
门户 http://127.0.0.1:5173    运营 http://127.0.0.1:5174    API :8080
假设需求 D-A260908：铝合金壳体精密制造 20000 件 AQL1.0 硬交期 22 天
假设分配：博锐 8000×3.00=24000；宏达 6000×4.00=24000；金盾 6000×4.00=24000；合计 72000
意向金 1000；工厂保证金=总报价×5%；买家保证金=预估总价×5%；佣金 1%"""


CASES = [
    ["AUTH-0001", "认证授权", "买家合法登录进入工作台", "功能", "高",
     "演示库已种子化，5173/8080 已启动",
     "URL=http://127.0.0.1:5173/login\n账号=13000000001 密码=123456",
     "1. 打开买家/工厂门户\n2. 输入账号密码点登录",
     "返回 token；role=BUYER；跳转 /buyer/home；工作台可见待办与企业名「杭州精工传动有限公司」。",
     "用户与企业-登录"],
    ["AUTH-0002", "认证授权", "工厂合法登录进入工作台", "功能", "高",
     "工厂 13000000002 存在",
     "账号=13000000002 密码=123456",
     "1. 登录工厂账号",
     "role=FACTORY，进入 /factory/home，可见浏览需求、我的报名、设备。",
     "用户与企业-登录"],
    ["AUTH-0003", "认证授权", "运营登录运营端", "功能", "高",
     "内置超级管理员",
     "URL=http://127.0.0.1:5174/login\n账号=admin 密码=admin123",
     "1. 打开运营端并登录",
     "进入管理控制中心，可见工作台、需求、资金、用户、日志。",
     "运营管理-登录"],
    ["AUTH-0004", "认证授权", "质检登录仅见质检台", "功能", "高",
     "质检账号已创建",
     "账号=inspect01 密码=123456",
     "1. 用质检账号登录 5174",
     "菜单裁剪为工单质检；不可见资金改余额、用户删除。",
     "运营管理-角色裁剪"],
    ["AUTH-0005", "认证授权", "错误密码不得登录且不泄露内部栈", "异常", "高",
     "买家账号存在",
     "账号=13000000001 密码=Wrong#999",
     "1. 输入错误密码提交",
     "提示登录失败；不写 token；停留登录页；页面无 SQL/堆栈。",
     "用户与企业-安全"],
    ["AUTH-0006", "认证授权", "空账号空密码前端拦截", "边界", "中",
     "打开登录页", "账号空 密码空", "1. 不填直接登录",
     "表单校验拦截，不发或即使发出也被拒绝。",
     "用户与企业-校验"],
    ["AUTH-0007", "认证授权", "买家不可访问工厂路由", "安全", "高",
     "买家已登录", "访问 /factory/home",
     "1. 地址栏改为工厂首页",
     "路由守卫重定向，不渲染工厂菜单。",
     "用户与企业-权限"],
    ["AUTH-0008", "认证授权", "工厂调用买家发布接口 403", "安全", "高",
     "工厂已登录", "POST /api/demand",
     "1. 用工厂 token 调发布需求",
     "403，demand 不新增。",
     "用户与企业-权限"],
    ["AUTH-0009", "认证授权", "未登录访问业务页跳转登录", "安全", "高",
     "已清空 sessionStorage", "访问 /buyer/demands",
     "1. 清除 token 刷新",
     "跳转 /login，接口 401。",
     "用户与企业-权限"],
    ["AUTH-0010", "认证授权", "买家注册成功并开户", "功能", "高",
     "使用未占用手机号",
     "手机=13900008801 密码=Abc123456\n企业=假设买家湖州精铸 信用代码=91330500MA8TEST801\n类型=BUYER 联系人=周启明 地址=湖州市吴兴区",
     "1. 登录页切注册并填齐提交",
     "可立即登录；account.balance=10000000.00，frozen=0；enterprise.type=BUYER。",
     "用户与企业-注册"],
    ["AUTH-0011", "认证授权", "重复手机号拒绝注册", "异常", "高",
     "13000000001 已存在", "手机=13000000001 其余合法",
     "1. 用已注册手机号注册",
     "提示手机号已注册，不新增 user/enterprise。",
     "用户与企业-校验"],
    ["AUTH-0012", "认证授权", "信用代码重复拒绝注册", "异常", "高",
     "91330110MA2K8B1X1A 已被占用",
     "新手机 13900008802 + 该信用代码",
     "1. 换手机号但信用代码重复提交",
     "提示信用代码已存在，事务回滚。",
     "用户与企业-校验"],
    ["DEM-0001", "需求发布", "发布 AQL 需求进入待审", "功能", "高",
     "买家已登录",
     "标题=假设-转向节壳体精铣 数量=20000\n材料=ADC12 AQL=1.0 意向期=1天 分期=1",
     "1. 打开发布需求按数据填写并提交",
     "demand.status=PENDING_AUDIT；inspectMode=AQL；列表可见。",
     "买家需求-发布"],
    ["DEM-0002", "需求发布", "全检必须填最低良率", "功能", "高",
     "发布页", "检验=FULL 先空良率再填 0.97",
     "1. 选全检不填良率提交应失败\n2. 填 0.97 再提交",
     "未填被拦截；成功后 inspectMode=FULL，aql 为空。",
     "买家需求-质量门槛"],
    ["DEM-0003", "需求发布", "必填缺失不得发布", "边界", "高",
     "发布页", "仅填标题",
     "1. 清空数量/材料/交期提交",
     "前端拦截；绕过后端也失败，不落库。",
     "买家需求-校验"],
    ["DEM-0004", "需求发布", "数量必须为正整数", "边界", "中",
     "发布页", "quantity=0",
     "1. 数量填 0 提交",
     "控件 min=1 或后端拒绝「数量无效」。",
     "买家需求-校验"],
    ["DEM-0005", "需求发布", "买家只能看自己的需求", "安全", "高",
     "买家 A 登录", "GET 需求列表",
     "1. 打开我的需求并尝试改 URL 看他人 id",
     "列表 tenant_id 均为当前买家；越权查看/修改失败。",
     "买家需求-隔离"],
    ["DEM-0006", "需求发布", "运营通过后进入意向期", "功能", "高",
     "需求 PENDING_AUDIT", "运营账号审核通过",
     "1. 运营打开待审需求点通过",
     "status=PUBLISHED；intentionEndAt 生效；工厂大厅可见。",
     "买家需求-审核"],
    ["DEM-0007", "需求发布", "运营退回必须填写原因", "异常", "中",
     "需求待审", "退回原因空 / 填「图纸公差未闭合」",
     "1. 空原因退回应失败\n2. 填写原因后退回",
     "空原因被拒；成功后 status=RETURNED，买家待办出现「退回修改」。",
     "买家需求-退回"],
    ["BID-0001", "工厂竞标", "意向期报名冻结 1000 元", "功能", "高",
     "需求 PUBLISHED；工厂余额≥1000",
     "工厂=13000000002 minQty=4000 maxQty=9000",
     "1. 浏览需求填承接区间报名",
     "quotation.status=INTENTION，intentionStatus=FROZEN；frozen+1000；流水 INTENTION/FREEZE。",
     "工厂竞标-意向金"],
    ["BID-0002", "工厂竞标", "余额不足禁止报名", "异常", "高",
     "将测试工厂可用余额临时改为 500（测后恢复）",
     "balance=500",
     "1. 尝试报名",
     "提示账户余额不足；不写 quotation，不写流水。",
     "工厂竞标-意向金"],
    ["BID-0003", "工厂竞标", "maxQty 不得超过需求数量", "边界", "高",
     "需求 quantity=20000", "minQty=1000 maxQty=30000",
     "1. 超限提交",
     "提示最大承接量不能超过需求数量 20000。",
     "工厂竞标-承接区间"],
    ["BID-0004", "工厂竞标", "minQty 不得大于 maxQty", "边界", "中",
     "意向期", "minQty=9000 maxQty=2000",
     "1. 反填区间提交",
     "提示最小量不能大于最大量。",
     "工厂竞标-承接区间"],
    ["BID-0005", "工厂竞标", "非意向期不可新报名", "异常", "中",
     "需求已 FACTORY_THINKING", "未报名工厂",
     "1. 思考期后尝试新报名",
     "拒绝报名，提示当前阶段不可报名。",
     "工厂竞标-状态机"],
    ["BID-0006", "工厂竞标", "思考期报价冻结 5% 保证金并退意向金", "功能", "高",
     "该厂已报名且需求=FACTORY_THINKING",
     "unitPrice=3.00 实施方案=五轴精铣+去毛刺\nmaxQty=9000 → 总报价 27000 → 保证金 1350",
     "1. 报名详情填报报价并确认",
     "deposit FROZEN=1350；意向金 1000 解冻回可用；通知买家「工厂已填报」。",
     "工厂竞标-保证金"],
    ["BID-0007", "工厂竞标", "思考期退出退意向金并记信用", "功能", "高",
     "FACTORY_THINKING 且未报价", "点取消报名",
     "1. 确认退出",
     "意向金解冻；不再进方案；credit_event 记 HONESTY_FACTORY_EXIT。",
     "工厂竞标-退出"],
    ["BID-0008", "工厂竞标", "无单价工厂不得进入方案", "异常", "高",
     "仅报名未填 unitPrice", "运营生成方案",
     "1. 检查候选工厂",
     "无有效单价的厂不进入锁定分配。",
     "方案编排-约束"],
    ["SOL-0001", "方案编排", "AI 方案从 AI1 连续编号", "功能", "高",
     "买家已交保证金，需求 SOLUTION_GENERATED",
     "锁定报价三厂可覆盖 20000 件",
     "1. 运营打开「审核 AI 推荐方案」\n2. 若历史为 AI3，点重新生成",
     "待审方案标题为方案 AI1（若有多套则为 AI1/AI2/AI3 连续），不得从 AI3 起跳。type 由后端按校验成功套数重排，不信任模型返回编号。",
     "方案编排-AI编号"],
    ["SOL-0002", "方案编排", "分配件数必须落在承接区间且合计等于需求", "功能", "高",
     "待审 AI 方案",
     "博锐 8000∈[4000,9000] 宏达 6000 金盾 6000 合计 20000",
     "1. 核对照表分配量",
     "每厂 quantity∈[minQty,maxQty]，合计=20000；小计=单价×件数。",
     "方案编排-约束"],
    ["SOL-0003", "方案编排", "区间无法凑齐不得下发", "异常", "高",
     "所有厂 max 之和 < 20000", "覆盖不足",
     "1. 触发生成或下发",
     "提示产能/承接不足，不能把件数对不上的方案 ACTIVE。",
     "方案编排-约束"],
    ["SOL-0004", "方案编排", "未勾选待审方案不能分次下发造成买家列表不同步", "异常", "高",
     "存在 2 套待审", "只下发其中一套后再下发另一套",
     "1. 分两次点下发",
     "系统要求勾选后一次下发；未选待审标 REJECTED；买家只见 ACTIVE。",
     "方案编排-审核"],
    ["SOL-0005", "方案编排", "运营一次下发后买家可见并可确认", "功能", "高",
     "勾选 AI1", "运营点「下发选中方案」",
     "1. 下发\n2. 买家打开需求",
     "买家待办「可选择方案」；可确认或自选后确认。",
     "方案编排-下发"],
    ["SOL-0006", "方案编排", "确认后分配不可再改", "功能", "高",
     "方案已下发", "买家确认",
     "1. 确认方案\n2. 再改件数",
     "生成 order + 分厂 contract 草稿；再次保存分配被拒绝。",
     "方案编排-确认"],
    ["SOL-0007", "方案编排", "AI 密钥失效不阻断主流程", "异常", "中",
     "无效 API Key", "重新生成 AI",
     "1. 触发 AI 生成",
     "运营收到生成失败通知；不 500；可重试。",
     "方案编排-AI"],
    ["CON-0001", "合同签署", "买家 48 小时内按厂上传合同", "功能", "高",
     "方案已确认，order.contractIssueEndAt=创建+48h",
     "为博锐上传 功能模块图.png <5MB",
     "1. 合同签署页对博锐点上传合同",
     "仅该行 attachmentId 有值；操作列显示「已上传，待统一下发」而不是「已下发」；倒计时仍走 48h 上传窗口。",
     "合同签署-上传"],
    ["CON-0002", "合同签署", "单份上传后对应工厂仍不可见合同", "功能", "高",
     "仅博锐已上传，买家未统一签字",
     "工厂 13000000002 打开报名详情",
     "1. 工厂查看合同文件\n2. 尝试签署",
     "提示买家尚未统一下发；无附件下载；actionKey=WAIT_ISSUE；待办不出现「待签署」。",
     "合同签署-可见性"],
    ["CON-0003", "合同签署", "未全部上传不得统一签字", "异常", "高",
     "三厂中仅一厂有附件", "买家点「签字并统一下发」",
     "1. 尝试提交签名",
     "提示请先分别上传「台州宏达…」等未上传工厂的合同。",
     "合同签署-统一下发"],
    ["CON-0004", "合同签署", "全部上传后买家签字即统一下发并关闭 48h 窗口", "功能", "高",
     "三厂附件齐全", "手写签名 + 勾选已阅读",
     "1. 提交签字并统一下发",
     "三份合同均写入 buyerSign；order.contractSignEndAt=现在+24h；contractIssueEndAt 清空；三厂同时收到「请签合同」；上传阶段结束。",
     "合同签署-统一下发"],
    ["CON-0005", "合同签署", "统一下发后工厂 24 小时内可签", "功能", "高",
     "已统一下发", "工厂阅读并签名",
     "1. 工厂打开合同并提交签名",
     "factorySign 落库；买家待办转为等待工厂或确认派单。",
     "合同签署-工厂签"],
    ["CON-0006", "合同签署", "统一下发后不可再更换合同文件", "异常", "高",
     "已统一下发", "买家再对博锐上传新文件",
     "1. 再次上传",
     "提示合同已统一下发，不能再修改。",
     "合同签署-不可更换"],
    ["CON-0007", "合同签署", "工厂逾期未签按规则处理", "异常", "中",
     "contractSignEndAt 已过且该厂未签", "定时 timeoutContractSign",
     "1. 等待调度或把截止时间改到过去后触发",
     "未签工厂保证金按规则划转赔偿买家；已签工厂不受误伤。",
     "合同签署-超时"],
    ["CON-0008", "合同签署", "买家 48h 未完成上传+签字则超时罚没", "异常", "中",
     "CREATED 且 issueEndAt 已过且尚未统一下发", "即使三厂附件已齐但未签字",
     "1. 触发 timeoutContractIssue",
     "按方案件数比重赔偿工厂；不能再补签下发。",
     "合同签署-超时"],
    ["CON-0009", "合同签署", "未全部工厂签字不得派单生产", "异常", "高",
     "买家已签，金盾未签", "点确认全部签署并开始派单",
     "1. 确认派单",
     "提示「嘉兴金盾…尚未完成双方签署」；不拆工单。",
     "合同签署-派单"],
    ["CON-0010", "合同签署", "全部签署后派单拆期生成工单", "功能", "高",
     "三厂均已双签", "买家确认派单",
     "1. 二次确认后派单",
     "work_stage 按分期拆出；需求进入生产；工厂可见待开工。",
     "合同签署-派单"],
    ["PRD-0001", "生产履约", "工厂上报进度且完成件数必须递增", "功能", "高",
     "工单 IN_PRODUCTION", "上次 1000，本次填 1000",
     "1. 上报与上次相同件数",
     "拒绝「必须大于上次上报」；改为 2500 并填说明后成功。",
     "生产履约-报工"],
    ["PRD-0002", "生产履约", "完成件数不得超过本段数量", "边界", "中",
     "本段 6000", "doneQty=7000",
     "1. 超交提交",
     "提示不能超过本段数量 6000。",
     "生产履约-报工"],
    ["PRD-0003", "生产履约", "未到分期开始不能报工", "异常", "中",
     "工单 WAITING_OPEN", "未到 periodStart",
     "1. 上报进度",
     "提示尚未到本期开始时间。",
     "生产履约-窗口"],
    ["PRD-0004", "生产履约", "非承接工厂不能报他厂工单", "安全", "高",
     "博锐工单", "宏达 token 调交付",
     "1. 越权提交",
     "403，数据不变。",
     "生产履约-隔离"],
    ["INS-0001", "质检", "买家支付质检费后进入待质检", "功能", "高",
     "PENDING_INSPECT_PAY payer=BUYER 实交=2500",
     "AQL 约 5 元/件 ×2500=12500（以系统计价为准）",
     "1. 买家点提交质检费",
     "扣买家余额入平台；status=PENDING_INSPECTION；质检台可见。",
     "质检管理-质检费"],
    ["INS-0002", "质检", "待付费工单对质检方不可见", "安全", "高",
     "仍 PENDING_INSPECT_PAY", "质检账号打开队列",
     "1. 刷新质检台",
     "该工单不出现。",
     "质检管理-可见性"],
    ["INS-0003", "质检", "抽检数少于 AQL 规定不得提交", "边界", "高",
     "PENDING_INSPECTION 实交 2500 AQL1.0 规定抽检 125",
     "sampleCount=80",
     "1. 质检提交",
     "提示抽检数不能少于规定的 125 件。",
     "质检管理-AQL"],
    ["INS-0004", "质检", "不合格件数不能超过抽检数", "边界", "中",
     "抽检 125", "关键不合格 80 + 一般 60",
     "1. 提交",
     "提示不合格件数不能超过抽检数。",
     "质检管理-校验"],
    ["INS-0005", "质检", "结论必须与数量/公差数据一致", "异常", "高",
     "数量不达标且公差不合格", "result 手选 PASS",
     "1. 提交审核",
     "提示是否合格须与检验数据一致。",
     "质检管理-一致性"],
    ["INS-0006", "质检", "质检填报后待运营审且不写信用事件", "功能", "高",
     "质检角色提交", "合法报告",
     "1. 质检员提交",
     "status=PENDING_REVIEW；通知运营；credit_event 不在此步插入。",
     "质检管理-填报"],
    ["INS-0007", "质检", "运营审核不合格写入长类型信用事件且不报截断", "功能", "高",
     "PENDING_REVIEW，结论 FAIL",
     "quantityOk=false toleranceOk=false",
     "1. 运营点提交审核",
     "不再出现 Data too long for column type；credit_event 写入 QUALITY_FAIL 与 HONESTY_INSPECT_FAIL；工单 status=FAIL。",
     "质检管理-信用字段"],
    ["INS-0008", "质检", "不合格后买家工作台出现待办", "功能", "高",
     "工单已 FAIL", "买家刷新工作台/通知",
     "1. 查看待办\n2. 点通知「去处理」",
     "待办 DECIDE_INSPECTION：「有 N 笔不合格工单待处理」；通知可跳转需求详情处理让步/返工/关闭。",
     "质检管理-待办"],
    ["INS-0009", "质检", "让步后改写应付并赔付", "功能", "中",
     "FAIL 且规则允许让步", "action=CONCESSION",
     "1. 买家确认让步",
     "payAmount 按公式更新；工厂保证金赔付买家；可支付托管。",
     "质检管理-让步"],
    ["INS-0010", "质检", "返工期限必须 12～72 小时", "边界", "中",
     "允许返工", "reworkHours=8 与 80",
     "1. 分别提交",
     "提示返工期限须为 12～72 小时。",
     "质检管理-返工"],
    ["INS-0011", "质检", "质量返工由工厂支付质检费", "功能", "中",
     "reworkKind=QUALITY 再交付", "payer=FACTORY",
     "1. 工厂支付质检费",
     "扣工厂余额；买家不再出现付费按钮。",
     "质检管理-质检费"],
    ["PAY-0001", "资金账户", "合格阶段余额托管成功", "功能", "高",
     "工单 PASS，买家余额充足",
     "应付=24000（博锐 8000×3.00）",
     "1. 买家确认余额支付托管",
     "买家余额-24000；平台+24000；工厂余额不变；escrow=HELD；流水 ESCROW OUT(买家)+IN(平台)。",
     "资金账户-托管"],
    ["PAY-0002", "资金账户", "运营总览平台托管不得因双分录显示 0", "功能", "高",
     "已发生上一笔托管 24000",
     "GET /api/fund/overview",
     "1. 运营打开资金总览刷新",
     "平台托管=平台账户 ESCROW IN−OUT，至少为 24000（若无结算）；不得把买家 ESCROW OUT 与平台 IN 轧差成 0。",
     "资金账户-总览口径"],
    ["PAY-0003", "资金账户", "点击四类 KPI 可看明细", "功能", "中",
     "运营资金页", "意向冻结/保证金冻结/平台托管/佣金累计",
     "1. 分别点击四张卡片",
     "弹出对应明细表，含时间、需求、主体、类型、方向、金额。托管明细只含平台租户 ESCROW。",
     "资金账户-明细"],
    ["PAY-0004", "资金账户", "需求流水能看到仅有 orderId 的托管分录", "功能", "中",
     "托管流水 demandId 原为空", "资金-需求流水",
     "1. 展开需求 D-A260908",
     "可见 ESCROW 行（通过 orderId 回填 demandId），不再被过滤。",
     "资金账户-需求流水"],
    ["PAY-0005", "资金账户", "余额不足拒绝托管", "异常", "高",
     "临时把买家可用改为 1000（测后恢复）", "应付=24000",
     "1. 确认支付",
     "提示余额不足；escrow 不变。",
     "资金账户-托管"],
    ["PAY-0006", "资金账户", "非 PASS 阶段不可支付", "异常", "高",
     "工单仍生产中", "直接 POST /order/stage/{id}/pay",
     "1. 调支付接口",
     "提示仅可收款阶段可以支付。",
     "资金账户-状态机"],
    ["PAY-0007", "资金账户", "重复支付幂等", "异常", "高",
     "该阶段已 HELD", "再次支付",
     "1. 二次请求",
     "拒绝重复入账；幂等号 ESCROW-OUT-{stageId}。",
     "资金账户-幂等"],
    ["PAY-0008", "资金账户", "他人民需求不可代付", "安全", "高",
     "买家 A 工单", "买家 B token",
     "1. 调 payStage",
     "403 只能支付自己的订单阶段款。",
     "资金账户-隔离"],
    ["PAY-0009", "资金账户", "未全部托管不能完工确认", "异常", "高",
     "仍有一期未 HELD", "买家点完工确认",
     "1. 验收",
     "提示仍有阶段未托管；订单不 COMPLETED。",
     "资金账户-结算"],
    ["PAY-0010", "资金账户", "完工确认按厂结算并分摊 1% 佣金", "功能", "高",
     "全部 PASS 且 HELD", "总额 72000 佣金 720",
     "1. 完工确认",
     "各厂到账=该厂工钱−分摊佣金；佣金入平台；剩余买家保证金退回；COMPLETED。",
     "资金账户-结算"],
    ["NTF-0001", "通知待办", "关键动作产生可跳转通知", "功能", "中",
     "完成报名或统一下发", "打开通知页",
     "1. 点击动作按钮",
     "跳到需求/报名详情，过期动作按钮自动隐藏。",
     "信用与通知"],
    ["NTF-0002", "通知待办", "操作写入 audit_log 且不可删", "功能", "中",
     "运营审核或买家支付后", "运营操作日志",
     "1. 按对象筛选",
     "可见操作人、动作、对象、时间。",
     "信用与通知"],
    ["OPS-0001", "运营管理", "资金流水类型只读可查", "功能", "中",
     "运营登录", "资金页",
     "1. 按企业/需求筛选",
     "可见 INTENTION/DEPOSIT/BUYER_DEPOSIT/ESCROW/COMMISSION/INSPECT_FEE；无直接改余额。",
     "运营管理-资金"],
    ["OPS-0002", "运营管理", "信用事件类型字段扩到 32 后历史库自动补丁", "功能", "中",
     "旧库 type=VARCHAR(16)", "重启后端 SchemaPatcher",
     "1. 查看 information_schema 列长度\n2. 再审一单 FAIL",
     "type 长度为 32；插入 HONESTY_INSPECT_FAIL 成功。",
     "运营管理-补丁"],
    ["GUI-0001", "通用GUI", "必填空提交拦截", "GUI", "中",
     "任意新建表单", "空表", "逐个必填留空提交", "红色校验，不成功提交。", "易用性"],
    ["GUI-0002", "通用GUI", "合同倒计时在上传期与签约期切换", "GUI", "中",
     "SOLUTION_SELECTED", "先看 issueEndAt，签字后再看 signEndAt",
     "1. 上传阶段观察剩余约 48h\n2. 统一下发后观察约 24h",
     "倒计时终点切换；文案从上传窗口变为工厂签约窗口。", "易用性"],
    ["GUI-0003", "通用GUI", "列表分页合法跳转", "GUI", "中",
     "记录数>一页", "页码=2 / 0 / 9999", "翻页与非法页", "第 2 页变化；非法页不 500。", "易用性"],
    ["GUI-0004", "通用GUI", "金额两位小数且状态标签与字典一致", "GUI", "低",
     "账户或需求列表", "任意金额/状态", "核对应付与标签", "¥ 两位小数；FAIL=待买家处理 等与 labels 一致。", "易用性"],
    ["CON-0011", "合同签署", "超 5MB 附件拒绝上传", "边界", "中",
     "方案已确认", "假设文件 功能模块图-6MB.png",
     "1. 选择超过 5MB 的文件上传",
     "前端或后端拒绝；attachmentId 不更新；不进入「已上传」。",
     "合同签署-上传"],
    ["CON-0012", "合同签署", "工厂不可越权签署他厂合同", "安全", "高",
     "已统一下发，博锐与宏达各有合同",
     "宏达 token + 博锐 contractId",
     "1. 调用工厂签署接口",
     "403；博锐 factorySign 仍为空。",
     "合同签署-隔离"],
    ["CON-0013", "合同签署", "买家未阅读勾选不得统一签字", "边界", "中",
     "三厂附件齐全", "签名有、勾选无",
     "1. 提交签字并统一下发",
     "前端拦截；后端即使收到也不写 buyerSign、不置 contractSignEndAt。",
     "合同签署-统一下发"],
    ["SOL-0008", "方案编排", "已下发 ACTIVE 方案重生成不得静默覆盖", "异常", "高",
     "买家已可见 ACTIVE 方案", "运营再点生成 AI",
     "1. 不撤回直接重生成",
     "拒绝覆盖或仅新增待审副本；买家当前 ACTIVE 方案不被悄悄替换。",
     "方案编排-约束"],
    ["PAY-0011", "资金账户", "质检员不能改账户余额", "安全", "高",
     "质检已登录", "任意改余额接口或 SQL 页面",
     "1. 尝试调整买家/工厂余额",
     "无入口；接口 403；account.balance 不变。",
     "资金账户-隔离"],
    ["NTF-0003", "通知待办", "过期任务动作按钮隐藏", "异常", "中",
     "合同 48h 已超时或需求已取消", "通知仍留在列表",
     "1. 打开历史通知点动作",
     "按钮隐藏或点击后提示任务已过期，不进入失效编辑页。",
     "信用与通知"],
    ["DEM-0008", "需求发布", "已发布门槛不可事后抬高", "异常", "高",
     "需求 PUBLISHED 或之后", "试图把 AQL 从 1.0 改为 0.65",
     "1. 买家编辑已发布需求门槛",
     "拒绝修改；工厂仍按原门槛履约。",
     "买家需求-质量门槛"],
    ["BID-0009", "工厂竞标", "同一工厂重复报名被拒绝", "异常", "中",
     "该厂已有 INTENTION", "再次提交相同区间",
     "1. 二次报名",
     "提示已报名；不重复冻结 1000 元。",
     "工厂竞标-意向金"],
]


def write_cases():
    wb = Workbook()
    info = wb.active
    info.title = "0-文档信息"
    info["B2"] = f"{TITLE} 测试用例"
    info["B2"].font = Font(name="微软雅黑", size=18, bold=True, color="1F3A5F")
    meta = [
        ["文档编号", "TC-第8组-001"],
        ["版本", VER],
        ["编制日期", TODAY],
        ["编制", MEMBERS],
        ["测试级别", "系统测试（功能 / 异常 / 边界 / 安全 / GUI）"],
        ["被测系统", f"{TITLE} 当前 master 演示版"],
        ["前端", "client-web :5173 ； admin-web :5174"],
        ["后端", "platform :8080"],
        ["数据库", "MySQL 8  dsh_platform"],
        ["通过准则", "步骤可复现；实际与预期一致；资金三账户可对账；无未关闭阻塞缺陷"],
        ["不通过准则", "主链路中断、托管显示 0、工厂提前看见合同、AI 从 AI3 起编、质检 FAIL 事务回滚、越权成功"],
        ["假设数据", DEMO.replace("\n", " | ")],
    ]
    for i, (k, v) in enumerate(meta, 4):
        info.cell(i, 2, k).font = Font(name="微软雅黑", bold=True, size=10)
        info.cell(i, 3, v).font = Font(name="微软雅黑", size=10)
        info.row_dimensions[i].height = 36 if i == 15 else 18
    info.column_dimensions["B"].width = 16
    info.column_dimensions["C"].width = 92

    write_rows(wb.create_sheet("1-测试策略"), ["项", "说明"], [
        ["测试范围", "登录权限、需求、报名报价、AI 方案编号与一次下发、合同 48h 上传/统一下发/24h 工厂签、报工、质检 FAIL 与待办、托管口径与 KPI 明细、结算佣金。鸿蒙端与真实支付宝列为观察项。"],
        ["测试方法", "手工系统测试；资金用例前后记录 balance/frozen；合同可见性必须买卖双方对照。"],
        ["测试数据", "演示账号见封面假设数据；资金异常允许临时改余额，测后恢复为买家 1000 万。"],
        ["费率时限", "意向金 1000；双端保证金 5%；佣金 1%；工厂/买家思考 24h；合同上传 48h；工厂签约 24h。"],
        ["本轮重点回归", "① AI 从 AI1 起编 ② 统一下发前工厂不可见 ③ 平台托管按平台租户统计 ④ credit_event.type≥32 ⑤ FAIL 待办"],
        ["通过标准", "高优先级 100% 通过；总体 ≥95%；无 Blocker。"],
    ], [16, 100])

    write_rows(wb.create_sheet("2-测试大纲"), ["序号", "一级", "二级", "三级", "缩写", "用例约"], [
        [1, "认证授权", "登录注册改密", "鉴权隔离", "AUTH", "12"],
        [2, "需求发布", "门槛/审核/隔离", "退回与门槛锁定", "DEM", "8"],
        [3, "工厂竞标", "意向金/保证金/区间", "退出与重复报名", "BID", "9"],
        [4, "方案编排", "AI1 起编/一次下发", "凑单与覆盖保护", "SOL", "8"],
        [5, "合同签署", "48h 上传/统一下发/24h 签", "可见性、超时、越权", "CON", "13"],
        [6, "生产履约", "递增报工/窗口", "越权", "PRD", "4"],
        [7, "质检", "AQL/FAIL/待办/信用字段", "返工", "INS", "11"],
        [8, "资金账户", "托管口径/明细/幂等", "结算与越权", "PAY", "11"],
        [9, "通知运营", "待办跳转/补丁/过期", "日志", "NTF/OPS", "5"],
        [10, "通用GUI", "倒计时/分页", "标签", "GUI", "4"],
    ], [8, 14, 28, 22, 12, 10])

    headers = ["用例编号", "所属模块", "用例标题", "用例类型", "优先级", "前置条件", "测试数据", "测试步骤", "预期结果", "追溯需求"]
    widths = [12, 12, 32, 8, 8, 30, 36, 36, 46, 18]
    func = [c for c in CASES if c[3] == "功能"]
    other = [c for c in CASES if c[3] not in ("功能", "GUI")]
    gui = [c for c in CASES if c[3] == "GUI"]
    write_rows(wb.create_sheet("3-功能测试用例"), headers, func, widths)
    write_rows(wb.create_sheet("4-异常边界安全用例"), headers, other, widths)
    write_rows(wb.create_sheet("5-通用GUI用例"), headers, gui, widths)
    write_rows(wb.create_sheet("6-版本记录"), ["版本", "日期", "作者", "说明"], [
        ["V1.0", "2026-08-20", MEMBERS, "按东软目录形成初稿"],
        ["V1.1", "2026-09-08", MEMBERS, "补主链路与资金用例"],
        [VER, TODAY, MEMBERS, "按当前实现重写：AI1 起编、合同统一下发可见性、托管口径、FAIL 待办、信用字段扩列，并加厚异常案例"],
    ], [10, 14, 28, 70])

    xlsx = ROOT / f"测试用例-{GROUP}-{LEADER}.xlsx"
    wb.save(xlsx)
    old = ROOT / f"测试用例-{GROUP}-{LEADER}.xls"
    if old.exists():
        old.unlink()
    return xlsx


def write_matrix():
    rows = [
        ["项目基础", "账号", "注册开户", "买家/工厂注册；买家默认 1000 万，工厂 100 万", "肖舒予", "M0", "已完成", "AUTH-0010~0012"],
        ["项目基础", "账号", "四角色登录", "买家/工厂/运营/质检分端登录", "肖舒予", "M0", "已完成", "AUTH-0001~0004"],
        ["项目基础", "权限", "路由与接口隔离", "跨角色 403，未登录 401", "肖舒予", "M0", "已完成", "AUTH-0007~0009"],
        ["需求竞标", "发布", "结构化需求+门槛", "AQL/全检互斥，门槛写死后不可抬高", "杜青桐", "M1", "已完成", "DEM-0001~0004"],
        ["需求竞标", "审核", "通过/退回", "退回必须原因；通过进意向期", "杜青桐", "M1", "已完成", "DEM-0006~0007"],
        ["需求竞标", "意向", "固定意向金 1000", "余额不足拒绝；区间合法", "李雨桐", "M1", "已完成", "BID-0001~0005"],
        ["需求竞标", "工厂思考", "报价锁 5% 保证金", "退意向金；退出记信用", "李雨桐", "M1", "已完成", "BID-0006~0008"],
        ["需求竞标", "买家思考", "预估总价 5% 买家保证金", "继续后异步生成 AI 方案", "李雨桐", "M1", "已完成", "SOL-0005"],
        ["方案合同", "AI 方案", "从 AI1 连续编号", "不信任模型 type；校验失败后重排", "马嘉祺", "M1", "已完成", "SOL-0001"],
        ["方案合同", "方案", "件数凑单约束", "∈区间且合计=需求；一次下发", "马嘉祺", "M1", "已完成", "SOL-0002~0005"],
        ["方案合同", "方案", "确认锁定", "确认后不可改，生成订单与合同草稿", "马嘉祺", "M1", "已完成", "SOL-0006"],
        ["方案合同", "合同上传", "48h 按厂上传", "下发前可换文件；工厂不可见", "杜青桐", "M2", "已完成", "CON-0001~0002"],
        ["方案合同", "统一下发", "签字即下发", "必须全部上传；关闭上传窗；开 24h 签约窗", "杜青桐", "M2", "已完成", "CON-0003~0006"],
        ["方案合同", "超时", "48h/24h 罚则", "未下发罚买家；未签罚对应工厂", "杜青桐", "M2", "已完成", "CON-0007~0008"],
        ["方案合同", "派单", "全签后拆工单", "缺一厂未签不可生产", "杜青桐", "M2", "已完成", "CON-0009~0010"],
        ["履约质检", "报工", "递增与窗口", "件数递增、不超段、未开窗拒绝", "肖舒予", "M2", "已完成", "PRD-0001~0004"],
        ["履约质检", "质检费", "谁触发谁付", "首次买家，质量返工工厂；未付质检不可见", "李雨桐", "M2", "已完成", "INS-0001~0002,INS-0011"],
        ["履约质检", "填报", "AQL 与一致性", "抽检下限、缺陷上限、结论与数据一致", "李雨桐", "M2", "已完成", "INS-0003~0006"],
        ["履约质检", "不合格", "信用+待办+处置", "type 扩 32；买家待办；让步/返工/关闭", "杜青桐", "M2", "已完成", "INS-0007~0010"],
        ["资金信用", "托管", "双分录+平台口径", "买家 OUT+平台 IN；总览只计平台净额", "马嘉祺", "M2", "已完成", "PAY-0001~0004"],
        ["资金信用", "托管", "门禁幂等", "非 PASS/余额不足/重复/越权拒绝", "马嘉祺", "M2", "已完成", "PAY-0005~0008"],
        ["资金信用", "结算", "佣金 1%", "按厂工钱占比分摊，退剩余保证金", "马嘉祺", "M2", "已完成", "PAY-0009~0010"],
        ["运营支撑", "资金明细", "四类 KPI 可点", "意向/保证金/托管/佣金细目", "马嘉祺", "M3", "已完成", "PAY-0003"],
        ["运营支撑", "通知", "过期动作隐藏", "请处理质检结果# 可跳转", "李雨桐", "M3", "已完成", "NTF-0001,INS-0008"],
        ["运营支撑", "补丁", "SchemaPatcher", "启动扩 credit_event.type", "杜青桐", "M3", "已完成", "OPS-0002"],
        ["运营支撑", "审计", "操作日志", "只增不删", "肖舒予", "M3", "已完成", "NTF-0002"],
        ["方案合同", "合同上传", "附件大小与签署勾选", "超 5MB 拒绝；未勾选已阅读不得下发", "杜青桐", "M2", "已完成", "CON-0011,CON-0013"],
        ["方案合同", "签署隔离", "工厂不可代签他厂", "token 与 contract.factory 必须一致", "杜青桐", "M2", "已完成", "CON-0012"],
        ["方案合同", "方案保护", "ACTIVE 不被静默覆盖", "已下发方案重生成须拒绝或另开待审", "马嘉祺", "M1", "已完成", "SOL-0008"],
        ["需求竞标", "发布", "门槛锁定", "已发布不得抬高 AQL/良率", "杜青桐", "M1", "已完成", "DEM-0008"],
        ["需求竞标", "意向", "防重复冻结", "同一工厂同一需求不可二次报名", "李雨桐", "M1", "已完成", "BID-0009"],
    ]
    wb = Workbook()
    cover_ws = wb.active
    cover_ws.title = "封面"
    cover_ws["B2"] = f"{TITLE} 需求跟踪矩阵（RTM）"
    cover_ws["B2"].font = Font(name="微软雅黑", size=18, bold=True, color="1F3A5F")
    cover_ws["B4"] = f"{GROUP}  {MEMBERS}  {VER}  {TODAY}"
    cover_ws["B5"] = "状态「已完成」= 当前代码已实现并纳入本轮系统测试。追溯列对应测试用例编号。"
    cover_ws["B7"] = "图例：已完成=实现+可测；进行中=部分实现；未开始=范围外。本表全部为已完成。"
    cover_ws.column_dimensions["B"].width = 100
    ws = wb.create_sheet("需求进度跟踪表")
    headers = ["NO.", "大分类", "中分类", "小功能", "详细说明（含现行业务规则）", "负责人", "里程碑", "状态", "追溯用例"]
    data = [[i] + r for i, r in enumerate(rows, 1)]
    write_rows(ws, headers, data, [6, 12, 12, 16, 58, 10, 10, 10, 22])
    for r in range(2, ws.max_row + 1):
        ws.cell(r, 8).fill = OKFILL
    write_rows(wb.create_sheet("填写说明"), ["项", "说明"], [
        ["覆盖率", f"矩阵 {len(rows)} 条需求点，对应用例 {len(CASES)} 条，高优先级必须可执行。"],
        ["与旧版差异", "合同由「一上传即下发」改为「全部上传+买家签字才统一下发」；AI 编号强制 AI1 起；托管 KPI 按平台租户。"],
        ["假设数据", "见测试用例封面。金额与件数用于课堂演示，与演示库数量级一致即可。"],
    ], [16, 90])
    xlsx = ROOT / f"需求矩阵-{GROUP}-{LEADER}.xlsx"
    wb.save(xlsx)
    old = ROOT / f"需求矩阵-{GROUP}-{LEADER}.xls"
    if old.exists():
        old.unlink()
    return xlsx


def write_srs():
    doc = setup_doc("需求规约")
    cover(doc, "需求规约", "依据 IEEE 830 裁剪 + 东软实训目录  · 对齐当前实现")
    add_h(doc, "修订记录", 1)
    add_table(doc, ["版本", "日期", "说明"], [
        ["V1.0", "2026-08-20", "按东软目录形成初稿"],
        ["V1.1", "2026-09-08", "补状态机、费率、门禁"],
        [VER, TODAY, "对齐当前代码：AI1 起编、合同统一下发、托管平台口径、不合格待办、信用字段扩列"],
    ], [3, 4, 10])

    add_h(doc, "1 引言", 1)
    add_h(doc, "1.1 目的", 2)
    add_p(doc, "本文规定智能制造云平台「做什么、不做什么、对错如何判定」。设计、编码、测试必须能回溯到本文条款。现场口头约定不得覆盖本文，除非走变更并同步矩阵与用例。", first=0.74, align="both")
    add_h(doc, "1.2 范围", 2)
    add_p(doc, "范围内：买家/工厂 Web、运营/质检 Web、Spring Boot 后端、MySQL 账本与状态机、演示种子数据。范围外必测：真实支付宝收单、银行存管、区块链。鸿蒙端复用同一套 API，作为扩展观察项。", align="both")
    add_h(doc, "1.3 背景", 2)
    add_p(doc, "中小工厂产能碎片化，大批量外协难以在一个平台闭环「比价—签约—质检—结算」。本系统用两阶段竞标降低报名门槛，用整品按件数分配避免按工序拆单扯皮，用独立质检与资金托管把「谁判定」和「谁出钱」分开。", align="both")
    add_h(doc, "1.4 参考资料", 2)
    add_p(doc, "《项目说明书》；application.yml 费率与时限；DemandStateMachine / ContractService / FundLedger / CreditScoring 现行实现；东软需求规约模板。")
    add_h(doc, "1.5 术语", 2)
    add_table(doc, ["术语", "现行定义"], [
        ["意向金", "报名冻结，固定 1000 元。报价成功或落选后退回。"],
        ["工厂保证金", "思考期报价后冻结，5%×（单价×maxQty）。"],
        ["买家保证金", "确认继续后按预估总价 5% 冻结，可抵尾款。"],
        ["统一下发", "全部合同上传完成后，买家一次签字；此前任何工厂都不可见附件。"],
        ["平台托管", "质检合格后买家划入平台的工费净额。总览只统计平台租户 ESCROW IN−OUT。"],
        ["AI1/AI2/AI3", "运营可见的方案业务编号。按校验成功的套数从 1 连续编号，禁止沿用模型返回的 AI3。"],
        ["整品分配", "一厂报一个单价，完成该需求全部工序，方案按件数拆。"],
    ], [4, 13])

    add_h(doc, "2 任务概述", 1)
    add_h(doc, "2.1 目标与成功标准", 2)
    add_p(doc, "目标：一条外协订单可在平台内从发布走到结算。成功标准：（1）四角色职责可独立完成；（2）买家/平台/工厂三账户在托管与结算时可对账；（3）统一下发前工厂不可见合同；（4）AI 方案从 AI1 起编；（5）不合格工单必有买家待办；（6）非法跳转、越权、重复入账均被拒绝。", align="both")
    add_h(doc, "2.2 用户特征", 2)
    add_table(doc, ["角色", "入口", "典型任务", "不能做"], [
        ["BUYER", "5173", "发需求、确认方案、上传并统一下发合同、付质检费/托管、处置不合格、完工", "改质检结论、看他租户单据"],
        ["FACTORY", "5173", "报名、报价、统一下发后签字、报工", "下发前看合同、直接取走托管款"],
        ["INSPECTION", "5174", "填质检单", "审结论、改余额、看未付费工单"],
        ["OPERATOR", "5174", "审需求/方案/质检、看资金明细与日志", "绕过账本直接改 balance"],
    ], [3, 2.2, 7, 5])

    add_h(doc, "2.3 主状态机", 2)
    add_p(doc, "实现以 DemandStateMachine 为准，禁止跳步。旧状态 THINKING/REVIEWING/LOCKING 仅兼容历史数据。")
    add_flow(doc, "图 2-1  需求主路径（结构）", [
        "开始",
        "  │  买家填单一品需求（图纸/门槛/意向期）",
        "  ▼",
        "PENDING_AUDIT ──运营退回──► RETURNED ──修改重提──► PENDING_AUDIT",
        "  │ 通过",
        "  ▼",
        "PUBLISHED（意向期，工厂锁 1000 意向金）",
        "  │ 到期 / 覆盖度满足",
        "  ▼",
        "FACTORY_THINKING（24h，报价锁 5% 保证金，退意向金）",
        "  │",
        "  ▼",
        "BUYER_THINKING（24h，继续则锁买家 5% 保证金；取消则退出）",
        "  │ 继续",
        "  ▼",
        "SOLUTION_GENERATED（AI 从 AI1 起编，运营一次下发）",
        "  │ 买家确认",
        "  ▼",
        "SOLUTION_SELECTED（48h 上传 → 签字统一下发 → 24h 工厂签 → 买家派单）",
        "  │",
        "  ▼",
        "IN_PRODUCTION → 质检 → 托管 → COMPLETED",
        "旁路：CANCELLED / FLOW_FAILED",
    ])

    add_h(doc, "3 需求规定", 1)
    add_h(doc, "3.1 一般性需求与配置", 2)
    add_table(doc, ["配置项", "默认值", "含义"], [
        ["dsh.fee.intention-fixed", "1000", "意向金（元）"],
        ["dsh.fee.deposit-rate", "0.05", "工厂保证金"],
        ["dsh.fee.buyer-deposit-rate", "0.05", "买家保证金"],
        ["dsh.fee.commission-rate", "0.01", "佣金"],
        ["factory / buyer-thinking-hours", "24 / 24", "思考期"],
        ["contract-issue-hours", "48", "上传+统一下发窗口"],
        ["contract-sign-hours", "24", "工厂签约窗口（自统一下发起）"],
    ], [6, 4, 7])

    add_h(doc, "3.2 功能性需求", 2)

    add_h(doc, "3.2.1 用户与企业", 3)
    add_table(doc, ["项", "内容"], [
        ["优先级", "高"],
        ["说明", "注册、登录、改密、按角色开户。JWT 24h。"],
        ["约束", "手机号 11 位唯一；信用代码 18 位唯一；密码 BCrypt；双层权限。"],
        ["验收", "错密无 token；买家打不开工厂页；重复键失败且无脏数据。"],
    ], [3, 14])

    add_h(doc, "3.2.2 需求发布与审核", 3)
    add_p(doc, "一单一品。质量门槛与分期一经发布不得事后抬高。AQL 与全检互斥：AQL 必选抽检水准；FULL 必填最低良率。", align="both")
    add_table(doc, ["字段", "约束"], [
        ["quantity", "正整数"],
        ["inspectMode", "AQL 或 FULL"],
        ["图纸附件", "必填，≤5MB"],
        ["intentionDays", "买家自定，到期由调度推进"],
    ], [4, 13])

    add_h(doc, "3.2.3 工厂报名与报价", 3)
    add_p(doc, "意向期只锁 1000 元；思考期报价锁保证金并退意向金。无单价不得进方案。", align="both")

    add_h(doc, "3.2.4 方案编排（含 AI 编号）", 3)
    add_p(doc, "运营审核后一次下发给买家。模型可能返回 AI3 或前两套 hydrate 失败，后端必须按成功套数重排为 AI1、AI2、AI3。已下发 ACTIVE 方案重生成时不得静默覆盖。", align="both")
    add_flow(doc, "图 3-1  AI 方案编号（结构）", [
        "模型 JSON schemes[]（可能带 type=AI3）",
        "  │  逐套 hydrate：工厂必须锁定报价，件数合计=需求",
        "  │  失败则丢弃该套，不占用序号",
        "  ▼",
        "成功套按插入顺序强制 type = AI + (index)",
        "  │",
        "  ▼",
        "PENDING_REVIEW → 运营勾选 → 一次 ACTIVE，其余 REJECTED",
    ])

    add_h(doc, "3.2.5 合同签署（统一下发）", 3)
    add_p(doc, "现行规则与旧版「上传即下发」不同，测试必须按本节执行。", align="both")
    add_table(doc, ["阶段", "时限", "谁可见附件", "结束条件"], [
        ["上传", "订单创建起 48h", "仅买家/运营", "全部附件齐全且买家统一签字"],
        ["工厂签约", "统一下发起 24h", "对应中标工厂可见本厂合同", "各厂签字"],
        ["派单", "全签后", "双方", "买家确认派单，拆 work_stage"],
    ], [3, 4, 5, 5])
    add_flow(doc, "图 3-2  合同统一下发（结构）", [
        "确认方案 → 为每厂插 DRAFT 合同，启动 48h",
        "  │  买家按厂上传（单份上传 ≠ 下发）",
        "  │  工厂此时：列表过滤掉未释放合同，签署接口拒绝",
        "  ▼",
        "全部有附件 → 买家一次签字",
        "  │  写 buyerSign；置 contractSignEndAt=+24h；清空上传窗",
        "  │  通知全部中标工厂「请签合同」",
        "  ▼",
        "工厂 24h 内签署 → 买家确认派单 → 生产",
    ])
    add_p(doc, "异常：缺附件不能签；已下发不能换文件；48h 未下发罚买家保证金；24h 未签罚对应工厂保证金。", align="both")

    add_h(doc, "3.2.6 生产报工", 3)
    add_p(doc, "仅本厂工单。完成件数必须大于上次且不超过本段数量。未到 periodStart 为待开启。实交件数是质检前置。", align="both")

    add_h(doc, "3.2.7 质检与不合格待办", 3)
    add_p(doc, "质检员只填报，运营出 PASS/FAIL。FAIL 时写入 QUALITY_FAIL 与 HONESTY_INSPECT_FAIL，故 credit_event.type 必须 VARCHAR(32)，否则整笔事务回滚，工单不会变 FAIL、买家也收不到待办。", align="both")
    add_table(doc, ["结论", "系统行为"], [
        ["PASS", "可支付托管；双方问卷"],
        ["FAIL", "工单待买家处理；待办 DECIDE_INSPECTION；通知「请处理质检结果#」可跳转"],
        ["让步", "改写应付，保证金按规则赔付"],
        ["返工", "12～72 小时；质量返工由工厂付质检费"],
        ["关闭", "关闭本段链路"],
    ], [3, 14])

    add_h(doc, "3.2.8 资金托管与总览", 3)
    add_p(doc, "托管写入一对分录：买家 ESCROW/OUT，平台 ESCROW/IN。总览若对全库 IN−OUT，买卖金额相抵恒为 0。需求规定：platformEscrow 只统计平台 tenant。保证金冻结含工厂 DEPOSIT 与买家 BUYER_DEPOSIT。四类 KPI 必须可点开细目。需求流水须先用 orderId 回填 demandId。", align="both")
    add_flow(doc, "图 3-3  托管记账（结构）", [
        "买家确认支付 24000",
        "  ├─ account.buyer.balance  −24000",
        "  ├─ account.platform.balance +24000",
        "  ├─ fund_flow ESCROW/OUT 24000 tenant=买家",
        "  └─ fund_flow ESCROW/IN  24000 tenant=平台",
        "总览 platformEscrow = Σ平台IN − Σ平台OUT",
    ])

    add_h(doc, "3.2.9 通知、信用与运营", 3)
    add_p(doc, "站内信按标题解析跳转，过期任务按钮隐藏。运营可建质检账号、只读流水、查看操作日志。SchemaPatcher 启动时扩列，兼容已有库。", align="both")

    add_h(doc, "3.3 安全性需求", 2)
    add_p(doc, "（S1）JWT 本机密钥。（S2）角色注解 + tenant 校验。（S3）密码哈希，上传分类型。（S4）资金与审核留痕、幂等键。（S5）演示默认走内部账本，避免未回调的「已付未入账」。", align="both")

    add_h(doc, "3.4 异常与边界目录（必须可测）", 2)
    add_p(doc, "下列情形不得导致 500、脏账或状态跳步。测试用例编号见矩阵追溯列。", align="both")
    add_table(doc, ["编号", "场景", "期望"], [
        ["E1", "错密 / 空表 / 未登录 / 跨角色", "无 token 或 401/403，无堆栈"],
        ["E2", "重复手机号或信用代码", "拒绝且事务回滚"],
        ["E3", "余额不足报名/托管", "不写 quotation / escrow 不变"],
        ["E4", "承接区间反填或超需求量", "拒绝冻结"],
        ["E5", "思考期后新报名、重复报名", "拒绝，不二次冻 1000"],
        ["E6", "覆盖不足或无单价厂", "方案不得 ACTIVE"],
        ["E7", "ACTIVE 方案被重生成", "不得静默覆盖买家当前方案"],
        ["E8", "缺附件、未勾选、未齐套签字", "不得统一下发"],
        ["E9", "单份上传后工厂拉详情", "不可见、不可签、无待办"],
        ["E10", "已下发再换文件 / 代签他厂", "拒绝"],
        ["E11", "48h 未下发 / 24h 未签", "只罚责任方保证金"],
        ["E12", "报工不递增、超段、未开窗、越权", "工单数量不变"],
        ["E13", "抽检不足、缺陷超抽检、结论打架", "不得提交"],
        ["E14", "信用 type 超长", "扩列后 FAIL 事务提交成功"],
        ["E15", "非 PASS / 重复支付 / 代付", "拒绝；幂等号不二次入账"],
        ["E16", "未全部托管点完工", "订单不 COMPLETED"],
        ["E17", "过期通知动作", "按钮隐藏"],
    ], [2.2, 7, 8])

    add_h(doc, "3.5 非功能", 2)
    add_p(doc, "性能：演示库单接口常规 <2s，资金写入必须同事务。可用性：状态与金额以库为准，刷新后与页面一致。兼容：Chrome 当前稳定版。可维护：费率与时限走配置；旧库靠 SchemaPatcher 扩列，禁止只改 schema.sql 不管存量。易用：合同倒计时必须随窗口切换，金额两位小数。", align="both")

    add_h(doc, "4 运行环境与接口", 1)
    add_p(doc, "JDK 21、MySQL 8、Vue 3、Element Plus。端口 8080/5173/5174。REST 前缀 /api。关键资源：/auth/login、/demand、/bidding、/solution/{id}/generate-ai、/order/{id}/contract/upload、/order/{id}/contract/buyer-sign、/order/stage/{id}/inspect、/order/stage/{id}/decision、/order/stage/{id}/pay、/fund/overview、/fund/details/{category}、/common/todos。", align="both")

    add_h(doc, "5 假设、遗留与验收", 1)
    add_p(doc, "假设演示账号与金额见测试用例封面。买家默认可用 10,000,000.00 元，工厂 1,000,000.00 元，便于走完整资金链而不依赖真实充值。假设订单「铝合金壳体精密制造」20000 件，三厂 8000/6000/6000，单价 3/4/4，总额 72000，佣金 720，仅用于对账与课堂演示，不代表真实成交价。", align="both")
    add_p(doc, "遗留：真实支付宝收单与回调对账、电子签章司法效力、鸿蒙端完整回归，不作为本次 Web 验收阻塞项，说明书中单列观察。验收：矩阵全部「已完成」项均可按对应用例走通；本轮五条回归（AI1 编号、统一下发可见性、托管平台口径、信用字段扩列、FAIL 待办）必须通过。", align="both")
    path = ROOT / f"需求规约-{GROUP}-{LEADER}.docx"
    doc.save(path)
    return path


def write_spec():
    doc = setup_doc("项目说明书")
    cover(doc, "项目说明书", "对应文件名：项目说明书-项目名称")
    add_h(doc, "1 项目概述", 1)
    add_h(doc, "1.1 项目背景", 2)
    add_p(doc, "制造业直接体现生产力水平。大量中小加工企业设备与订单分散，买家难以用统一规则完成外协协同。本项目在云端把买家、工厂、独立质检与平台运营连成同一条履约链路：需求可发布、工厂可竞标、方案可审核、合同可验证、资金可托管、质量有第三方结论。", first=0.74, align="both")
    add_h(doc, "1.2 项目名称与目标", 2)
    add_p(doc, "项目名称：智能制造云平台。目标不是再做一套报工表，而是把「分散的制造协作」变成「可验证的履约」：每一步有状态，每一笔钱有账本，每一次不合格有人处理。", align="both")
    add_h(doc, "1.3 关键技术", 2)
    add_table(doc, ["层次", "技术"], [
        ["前端", "Vue 3、Vite、Vue Router、Element Plus、Axios、GSAP"],
        ["管理端", "独立 admin-web，与门户同设计语言"],
        ["后端", "Spring Boot 3、Spring Security、JWT、MyBatis-Plus"],
        ["数据", "MySQL 8；启动期 SchemaPatcher 兼容旧库"],
        ["资金", "FundLedger 唯一写入口，幂等号防重"],
        ["方案", "规则生成 + DeepSeek JSON 方案，编号后端重排"],
        ["扩展", "鸿蒙 ArkTS 客户端复用 REST"],
    ], [3, 14])

    add_h(doc, "1.4 角色与端", 2)
    add_table(doc, ["角色", "端", "职责"], [
        ["买家 BUYER", "client-web :5173", "发布需求、确认方案、上传并统一下发合同、支付、处置不合格、完工"],
        ["工厂 FACTORY", "client-web :5173", "报名报价、下发后签约、报工交付"],
        ["质检 INSPECTION", "admin-web :5174", "填写质检单，不出资金结论"],
        ["运营 / 超管", "admin-web :5174", "审核、下发方案、审质检、资金明细、用户与日志"],
    ], [4, 4, 9])

    add_h(doc, "2 需求分析", 1)
    add_p(doc, "原始东软订单/计划/工单模型被落实为「需求—竞标—方案—合同—工单—质检—托管—结算」。计划不再由运营手工拆单，而由买家确认的方案件数与分期规则自动拆 work_stage。", align="both")
    add_h(doc, "2.1 功能结构", 2)
    add_flow(doc, "图 2-1  系统功能结构", [
        "智能制造云平台",
        "├─ 门户（买家/工厂）",
        "│   ├─ 注册登录 / 企业账户",
        "│   ├─ 需求发布与详情",
        "│   ├─ 报名报价 / 方案确认",
        "│   ├─ 合同上传与统一下发 / 工厂签署",
        "│   └─ 报工、质检处置、托管、完工",
        "├─ 运营指挥台",
        "│   ├─ 需求审核 / AI 方案审核下发",
        "│   ├─ 质检审核",
        "│   ├─ 资金总览与四类明细",
        "│   └─ 用户、设备档案、操作日志",
        "└─ 领域服务",
        "    ├─ 状态机与时限调度",
        "    ├─ 资金账本 FundLedger",
        "    └─ 信用事件 CreditScoring",
    ])

    add_h(doc, "2.2 端到端业务流程", 2)
    add_flow(doc, "图 2-2  主业务流程图", [
        "开始 → 注册/登录",
        " → 买家发布需求 → 运营审核",
        " → 工厂意向报名（锁 1000）",
        " → 工厂思考期报价（锁 5% 保证金）",
        " → 买家思考期交 5% 保证金",
        " → AI/规则出方案（AI1 起）→ 运营一次下发",
        " → 买家确认 → 48h 按厂上传合同",
        " → 买家签字统一下发（工厂此前不可见）",
        " → 工厂 24h 签署 → 买家派单",
        " → 生产报工 → 付质检费 → 质检填报 → 运营审核",
        "    ├─ PASS → 托管工费 → 问卷 → 全部完成后结算（佣金 1%）",
        "    └─ FAIL → 买家待办：让步 / 返工 / 关闭",
        " → 结束",
    ])

    add_h(doc, "3 总体设计", 1)
    add_h(doc, "3.1 架构", 2)
    add_p(doc, "前后端分离。Vite 开发服务器把 /api 代理到 8080。后端按领域拆服务：Demand / Bidding / Solution / Contract / Order / Fund / Credit / Notify，禁止在业务服务里直接插资金流水。", align="both")
    add_h(doc, "3.2 核心数据", 2)
    add_table(doc, ["表", "作用"], [
        ["enterprise / sys_user / account", "租户、账号、余额与冻结"],
        ["demand / process / quotation", "需求、工序、报名报价"],
        ["solution / order / contract / work_stage", "方案、订单、一厂一合同、分期工单"],
        ["inspection / survey / credit_event", "质检、互评、信用"],
        ["fund_flow / notify / audit_log", "账本、站内信、审计"],
    ], [6, 11])
    add_h(doc, "3.3 关键设计决策", 2)
    add_table(doc, ["问题", "决策", "原因"], [
        ["合同何时对工厂可见", "买家统一签字后", "避免一厂提前拿到文本、其余厂尚未齐套"],
        ["AI 方案编号", "后端按成功套数重排", "模型常返回 AI3 或前两套校验失败"],
        ["托管总览", "只计平台租户", "买卖双分录全库轧差恒为 0"],
        ["质检与放款", "结论与资金分离", "质检只判断，运营审核，平台执行已确认结果"],
        ["信用 type 长度", "VARCHAR(32)+启动补丁", "HONESTY_INSPECT_FAIL 等超过 16"],
    ], [4, 4, 9])

    add_h(doc, "4 详细设计要点", 1)
    add_p(doc, "合同释放条件：attachmentId 非空 ∧ buyerSign 非空 ∧ order.contractSignEndAt 非空。工厂列表、getMine、factorySign、待办 SIGN_CONTRACT 均按此过滤。资金明细接口 GET /api/fund/details/{INTENTION|DEPOSIT|ESCROW|COMMISSION}。不合格待办查询 work_stage.status=FAIL。", align="both")
    add_h(doc, "4.1 合同域", 2)
    add_p(doc, "确认方案时按中标工厂各插一条 DRAFT。上传只写附件，不打开工厂可见性。买家签字走 buyerSignAll：校验齐套后批量写 buyerSign，再置 contractSignEndAt=+24h，并清空上传截止。timeoutContractIssue 只处理尚未释放的订单；timeoutContractSign 只罚未签工厂。", align="both")
    add_h(doc, "4.2 方案编号", 2)
    add_p(doc, "AiSolutionGenerator 对模型返回的 type 不信任。hydrate 失败的套直接丢弃。成功套按插入顺序强制 type=\"AI\"+序号，保证运营审核页从 AI1 起编。", align="both")
    add_h(doc, "4.3 资金口径", 2)
    add_p(doc, "escrowStage 写买家 ESCROW/OUT 与平台 ESCROW/IN。overview.platformEscrow 只聚合平台 tenant。保证金冻结含 DEPOSIT 与 BUYER_DEPOSIT。byDemand 先用 orderId 回填 demandId，避免托管行从需求流水消失。", align="both")
    add_h(doc, "4.4 质检与待办", 2)
    add_p(doc, "approveInspect 在同一事务写工单 FAIL、QUALITY_FAIL、HONESTY_INSPECT_FAIL。credit_event.type 必须 ≥32，否则整笔回滚，买家待办也不会出现。TodoService 统计 FAIL；NotifyActionResolver 识别「请处理质检结果#」。", align="both")

    add_h(doc, "5 实现与环境", 1)
    add_table(doc, ["项", "内容"], [
        ["演示买家", "13000000001 / 123456  杭州精工传动有限公司"],
        ["演示工厂", "13000000002～05 / 123456"],
        ["运营", "admin / admin123"],
        ["假设订单", "铝合金壳体 20000 件；博锐 8000×3；宏达 6000×4；金盾 6000×4；总额 72000"],
        ["启动", "后端 8080，门户 5173，运营 5174，MySQL 本地 dsh_platform"],
    ], [3, 14])

    add_h(doc, "6 测试与结论", 1)
    add_p(doc, "系统测试用例见《测试用例-第8组-杜青桐》。本轮补强异常：越权、余额不足、区间非法、统一下发前可见性、AI 编号、托管口径、信用截断、FAIL 待办、幂等支付。结论：Web 主链路可演示、可对账、可回归。", align="both")

    add_h(doc, "7 小组分工", 1)
    add_table(doc, ["成员", "主要工作"], [
        ["杜青桐", "状态机与合同统一下发、质检不合格待办、文档统稿"],
        ["李雨桐", "竞标/保证金/质检费与通知跳转"],
        ["马嘉祺", "方案编号、资金账本口径与运营资金明细"],
        ["肖舒予", "认证权限、报工隔离、页面与账号种子"],
    ], [3, 14])
    path_docx = ROOT / "项目说明书-项目名称.docx"
    doc.save(path_docx)
    # 覆盖原 .doc：写入 OOXML 副本，WPS/Word 可打开；同时保留 docx
    path_doc = ROOT / "项目说明书-项目名称.doc"
    doc.save(path_doc)
    return path_docx


def write_intern():
    doc = Document()
    s = doc.sections[0]
    s.top_margin = Cm(2.5)
    s.bottom_margin = Cm(2.5)
    s.left_margin = Cm(2.8)
    s.right_margin = Cm(2.6)
    add_p(doc, "本科生毕业专业实训报告", 22, True, "center", 20, "黑体")
    add_p(doc, "项目名称：智能制造云平台", 14, True, "center", 18, "黑体")
    add_table(doc, ["项", "内容"], [
        ["学院", "人工智能学院（数据学院）"],
        ["年级专业", "2024 级数据科学与大数据技术"],
        ["学号姓名", "2024082101  杜青桐"],
        ["学生班级", "数据 2401 班"],
        ["企业导师", "张伟"],
        ["校内教师", "李娜、石宇欣"],
        ["实训时间", "2026 年 7 月 27 日—2026 年 9 月 10 日"],
        ["实训小组", "第 8 组  杜青桐、李雨桐、马嘉祺、肖舒予"],
    ], [4, 12])
    add_p(doc, "同组成员（假设学号，提交前可按学籍替换）：", 11, True)
    add_table(doc, ["姓名", "学号", "分工", "备注"], [
        ["杜青桐", "2024082101", "状态机、合同统一下发、质检待办、文档统稿", "组长"],
        ["李雨桐", "2024082102", "竞标保证金、质检费、通知跳转", "组员"],
        ["马嘉祺", "2024082103", "AI 编号、资金口径、运营明细", "组员"],
        ["肖舒予", "2024082104", "认证权限、报工隔离、页面与种子数据", "组员"],
    ], [3, 4, 7, 3])

    add_h(doc, "一、前言与知识", 1)
    add_h(doc, "1.1 实习收获（知识、能力、素养）", 2)
    add_p(doc, "知识方面，这次实训把离散的课程知识点重新焊成一条能跑通的业务链。Java 不再停留在类与接口作业，而是用来表达 Demand、Contract、WorkStage 这些有生命周期的对象；关系库不再是「建表拿分」，而是要同时回答租户隔离、幂等流水和状态不可跳步。Spring Boot 的分层让我看清：Controller 只做鉴权入口，领域规则必须下沉到 Service 与 Ledger，否则资金会在多个类里被各写一遍。Vue 3 的组合式 API 用来承载长页面状态，例如合同列表的倒计时要从「上传 48 小时」切到「签约 24 小时」，切错了用户会按错误窗口行动。Axios 与 JWT 则让我理解前后端分离不是少写 JSP，而是契约：401 必须回登录，403 必须是越权，业务失败必须是中文原因而不是堆栈。更重要的是一组制造领域知识：意向金与保证金为什么不能合成一笔，质检为什么不能直接放款，合同为什么必须整单齐套后才对工厂可见。这些知识如果只背定义，考试能写；一旦写进代码，错一个条件就会在路演当场穿帮。", first=0.74, align="both")
    add_p(doc, "能力方面，最大的进步是把模糊句子翻译成可验收规则。产品说「工厂不能提前看到合同」，如果只改按钮文案，工厂仍可能打详情接口把附件拉下来。正确做法是同一条释放条件同时卡住列表、详情和签署：必须有附件、必须有买家签字、订单上必须已经写下签约截止时间。产品说「托管不该是 0」，如果只在前端写死一个数，账本一对双分录仍然会把全库 IN 与 OUT 抵消。正确做法是总览按平台租户聚合。能力的第二层是用异常用例倒逼主路径。余额不足、承接区间反填、抽检数不够、结论与数据打架、重复支付、用别人的 token 代付——这些看起来像找茬，其实是在保护账本。能力的第三层是联调与对账：每一次托管都要同时看买家余额、平台余额、工厂余额和 fund_flow 四张证据，而不是看页面提示「成功」就结束。", first=0.74, align="both")
    add_p(doc, "素养方面，工业软件首先要对得起责任边界。质检只判断、运营来审核、平台只执行已经确认的结果；谁也不能因为「方便演示」去直接改 account.balance。文档、代码、用例必须使用同一套名词，否则说明书写「运营审合同」，代码却是买家统一下发，评委一问就裂。小组约定交叉执行测试，编制人不得只测自己模块。对真实支付宝、电子签章法律效力和鸿蒙端完整回归，选择在说明书里写成观察项，而不是为了材料好看写成「已完成」。也学会了在截止日前做减法：先守住主链和账本，再谈动画和口号。这些素养比多写一个页面更接近以后入职要面对的约束。", first=0.74, align="both")

    add_h(doc, "1.2 项目背景与意义", 2)
    add_p(doc, "制造业仍是国民经济的底座，但大量中小加工企业的设备、工艺和订单是散的。买家要做一笔两万件的壳体外协，往往要同时问三家工厂：谁能接、谁报价真、谁交得出、谁质检过、谁最后拿得到钱。传统做法把这些环节摊在微信群、电子表格和口头承诺里，门槛可以事后抬高，合同可以只给其中一厂看，质检结论可以和放款绑在同一个人手里，出了问题没有公共账本。智能制造云平台的意义，是把买家、工厂、独立质检和平台运营放进同一条可验证的履约链：需求带质量门槛发布，工厂先交象征性意向金再交保证金报价，运营审核方案，合同必须齐套后才统一下发，工费先入平台再按已确认的质检结果结算。它不是再做一个报工表，而是让协作可核对、资金可追溯、异常有人处理。", first=0.74, align="both")

    add_h(doc, "二、项目——智能制造云平台", 1)
    add_h(doc, "2.1 项目介绍", 2)
    add_p(doc, "系统采用前后端分离：买家/工厂门户与运营/质检端均为 Vue 3 + Element Plus，后端为 Spring Boot 3 + Spring Security + JWT + MyBatis-Plus，数据为 MySQL 8，库名 dsh_platform。核心模块包括用户与企业开户、需求发布与审核、两阶段竞标（意向金 1000、工厂/买家保证金各 5%）、规则与 DeepSeek AI 方案（编号由后端从 AI1 重排）、合同 48 小时上传后统一下发、工厂 24 小时签署、分期工单报工、独立质检与不合格处置、工费托管与 1% 佣金结算、站内通知与待办。演示假设订单为「铝合金壳体精密制造」20000 件，三厂分配博锐 8000×3.00、宏达 6000×4.00、金盾 6000×4.00，总额 72000 元，佣金 720 元，专门用于三账户对账与课堂演示，不代表真实成交。", first=0.74, align="both")
    add_flow(doc, "图 2-1  项目业务主链（结构）", [
        "发布需求 → 运营审核 → 意向报名(1000) → 工厂报价(5%)",
        " → 买家保证金(5%) → AI1 起编方案 → 一次下发 → 买家确认",
        " → 48h 上传合同 → 签字统一下发 → 24h 工厂签 → 派单生产",
        " → 报工 → 质检 → PASS 托管 / FAIL 买家处置 → 结算(佣金1%)",
    ])

    add_h(doc, "2.2 项目实施", 2)
    add_h(doc, "2.2.1 人员分工", 3)
    add_p(doc, "本人杜青桐负责需求状态机与合同域：把「上传即可见」改成「齐套签字后统一下发」，补齐工厂待办过滤和 48h/24h 超时口径，并处理不合格订单的买家待办与通知跳转。同时负责五份交付文档与现行代码对齐。李雨桐负责竞标资金与质检费支付方切换；马嘉祺负责方案编号重排和资金总览口径；肖舒予负责认证隔离与报工校验。每日站会同步接口契约，避免前后端口径漂移。", first=0.74, align="both")

    add_h(doc, "2.2.2 项目过程", 3)
    add_p(doc, "2.2.2.1 需求分析。对照东软原始「订单—计划—工单」描述，用逆向思维还原真实外协痛点：门槛必须事先写死，否则质检时双方会争议；报名必须有成本，否则工厂刷单；合同必须齐套下发，否则工厂之间信息不对称。据此确认业务对象、功能清单和两周可实施边界，并把思考期固定为 24 小时、意向金固定 1000 元写入配置即需求。", first=0.74, align="both")
    add_p(doc, "2.2.2.2 概要设计。抽取状态机、账本、合同释放三个不变点。状态机禁止跳步；FundLedger 是资金唯一写入入口；合同对工厂可见必须同时满足附件、买家签字和签约窗口。数据库按租户隔离，方案 type 与主键 id 分离。界面按角色裁剪菜单。", first=0.74, align="both")
    add_flow(doc, "图 2-2  合同释放判定（结构）", [
        "工厂请求合同",
        "  ├─ 无本厂草稿 → 403",
        "  ├─ 无附件或买家未签或未置签约截止 → 仍不可见",
        "  └─ 三者皆满足 → 可见并允许签署",
    ])
    add_p(doc, "2.2.2.3 编码。后端按领域服务拆分，前端按买家需求详情与工厂报名详情承载主操作。本人实现的关键点包括：buyerSignAll 后才打开 24h 窗口；listByOrder 对工厂过滤未释放合同；TodoService 统计 FAIL 工单；NotifyActionResolver 识别「请处理质检结果#」。联调时用演示账号走完整链路，而不是只测单接口 200。", first=0.74, align="both")
    add_p(doc, "2.2.2.4 测试。按功能、异常、边界、安全、GUI 五类编写用例。重点异常：工厂在单份上传后不应看到合同；AI 不得从 AI3 起编；运营审核 FAIL 不得因 type 超长回滚；托管总览不得显示 0。资金类用例执行前后记录余额，测后恢复。交叉执行，编制人不得只测自己模块。", first=0.74, align="both")
    add_p(doc, "2.2.2.5 文档。说明书、规约、矩阵、用例与实习报告使用同一套术语和同一组假设数据，流程图先写清结构再填规则，避免「图是旧流程、字是新流程」。", first=0.74, align="both")

    add_h(doc, "2.2.3 项目效果", 3)
    add_p(doc, "当前 Web 主链路可演示：四角色登录、需求从待审到完成、合同齐套下发、不合格可待办、资金总览可点明细。假设订单总额 72000 元，托管后平台净额与买家减少额一致，工厂余额在托管瞬间不变，结算时再按工钱减 1% 佣金入账。页面采用统一的蓝紫玻璃质感，工作台与登录页服务履约叙事而不是堆砌控件。效果如图 2-3 所示的结构（实现界面见系统截图附件，提交时按学院要求插入）。", first=0.74, align="both")
    add_flow(doc, "图 2-3  路演界面结构（插图位置）", [
        "[图 2-3  买家需求详情-合同签署区]",
        "说明：倒计时、按厂上传、底部「签字并统一下发」。",
        "[图 2-4  运营资金总览-四类 KPI]",
        "说明：平台托管为平台净额，卡片可点开细目。",
    ])
    add_p(doc, "插图规范：采用嵌入而非浮动；图题居中，中文宋体、英文 Times New Roman；图 2-3 与正文之间空 2 个半角空格的习惯按学院模板执行。若打印预览裁切，改为嵌入并限制宽度。", first=0.74, align="both")

    add_h(doc, "三、实训心得体会", 1)
    add_p(doc, "回看这一个多月，最大的收获不是又多会了几个框架注解，而是承认：工业协作软件能不能上桌，往往不取决于功能清单写了多少条，而取决于几条不起眼的约束有没有被守住。合同多上传一个文件，工厂就能提前看见，这不是按钮文案没改好，是释放条件没有同时卡住列表、详情和签署。托管数字显示 0，不是买家没付钱，是总览把买家的 ESCROW/OUT 和平台的 ESCROW/IN 放进同一个池子轧差。质检不合格后买家工作台空空，表面是提醒漏了，根子可能是 credit_event.type 只有 16 个字符，HONESTY_INSPECT_FAIL 一插入就把整笔事务回滚，工单甚至不会变成 FAIL。这三件事让我形成一个习惯：先画「谁在什么状态能看见什么、钱从哪个账户到哪个账户」，再打开编辑器写页面。", first=0.74, align="both")
    add_p(doc, "第二个体会是文档、代码、用例必须共用一套名词。我们曾经在说明书里还写着「运营确认双签」，代码却已经改成买家统一下发；路演时评委按说明书提问，现场就要改口。后来小组约定：状态名以 DemandStateMachine 为准，费率以 application.yml 为准，可见性以 ContractService.releasedToFactories 为准，任何口头约定都要回写规约和矩阵。测试也不再是「点一遍主路径截图」，而是故意做反例：余额改成 500 再报名、区间反填、抽检少填、用宏达的 token 签博锐的合同、对已托管工单再点一次支付。这些案例看起来像找茬，实际上是在保护账本和状态机。交叉执行也很必要——编制人测自己模块，很容易把「我知道该怎么点」当成系统真的防住了。", first=0.74, align="both")
    add_p(doc, "第三个体会是协作与取舍。四个人同时改同一条主链，接口契约一天变两次就会全员返工。我们用每日短会只同步三件事：今天合入了哪条规则、哪张表的口径变了、哪条用例要重跑。截止日前也学会做减法：真实支付宝、电子签章司法效力、鸿蒙端完整回归，明知重要，但没有回调对账就不敢写成「已完成」。说明书里把它们列为观察项，比材料好看更负责。假设订单 20000 件、总额 72000、学号 2024082101 起编，都是课堂可对账数据；若与学籍或现场库不一致，以学院登记和演示库为准替换，不在报告里假装那是生产成交。", first=0.74, align="both")
    add_p(doc, "不足同样清楚。我对调度超时的现场验证仍偏「改截止时间再触发」，而不是等足 48 小时；对 AI 方案只验证了编号重排，没有穷举模型各种残缺 JSON；页面动效服务了观感，但在弱网下的加载态还可以更克制。以后若继续做这类系统，会把「幂等号、租户、状态、账本」四张清单做成检查表，每合入一个接口就勾一次。整体上，这次实训把「会写 CRUD」推进到了「会守住一条履约链」。这比多做一个页面更接近以后入职要面对的约束：责任有边界，钱不能抄近路，细节必须经得起别人按文档来问。", first=0.74, align="both")

    path = ROOT / "实习报告-学号姓名.docx"
    doc.save(path)
    alias = ROOT / "实习报告-2024082101-杜青桐.docx"
    doc.save(alias)
    return path


def main():
    cases = write_cases()
    matrix = write_matrix()
    srs = write_srs()
    spec = write_spec()
    intern = write_intern()
    print("OK")
    for p in (cases, matrix, srs, spec, intern):
        print(p)


if __name__ == "__main__":
    main()
